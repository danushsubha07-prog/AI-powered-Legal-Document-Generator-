from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF

from utils.text_utils import sanitize_text


ASSET_DIR = Path(__file__).resolve().parents[1] / "assets"
LOGO_PATH = ASSET_DIR / "logo.png"


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def _set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def format_docx(text: str, doc_type: str, parties: str = "", terms: str = "") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO_PATH.exists():
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.8))
    else:
        p.add_run("LegalEase")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("AI-assisted legal document draft")
    sr.italic = True
    sr.font.size = Pt(9)

    if parties or terms:
        table = document.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        _set_cell_text(table.rows[0].cells[0], "Field", True)
        _set_cell_text(table.rows[0].cells[1], "Provided Information", True)
        if parties:
            row = table.add_row().cells
            _set_cell_text(row[0], "Parties", True)
            _set_cell_text(row[1], parties)
        if terms:
            row = table.add_row().cells
            _set_cell_text(row[0], "Terms", True)
            _set_cell_text(row[1], terms)

        document.add_paragraph()

    for block in sanitize_text(text).split("\n"):
        line = block.strip()
        if not line:
            document.add_paragraph()
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        if line.isupper() and len(line) < 100:
            r = p.add_run(line)
            r.bold = True
        else:
            r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.add_run("LegalEase • AI-assisted draft • Review before use\n")
    _add_page_number(fp)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class BrandedPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.doc_type = doc_type
        self.set_margins(18, 24, 18)
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=85, y=7, w=40)
        else:
            self.set_font("Times", "B", 13)
            self.cell(0, 8, "LegalEase", align="C")
        self.ln(13)
        self.set_draw_color(100, 100, 100)
        self.line(18, 22, 192, 22)
        self.ln(5)

    def footer(self):
        self.set_y(-14)
        self.set_font("Times", "I", 8)
        self.cell(0, 5, "LegalEase - AI-assisted draft - Review before use", align="C")
        self.ln(4)
        self.cell(0, 4, f"Page {self.page_no()}", align="C")


def format_pdf(text: str, doc_type: str, parties: str = "", terms: str = "") -> bytes:
    pdf = BrandedPDF(doc_type)
    pdf.add_page()
    pdf.set_title(doc_type)

    pdf.set_font("Times", "B", 16)
    pdf.cell(0, 9, doc_type.upper(), align="C")
    pdf.ln(10)

    if parties:
        pdf.set_font("Times", "B", 10)
        pdf.cell(0, 6, "PARTIES", ln=True)
        pdf.set_font("Times", size=10)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 5, sanitize_text(parties))
        pdf.ln(2)

    if terms:
        pdf.set_font("Times", "B", 10)
        pdf.cell(0, 6, "KEY TERMS", ln=True)
        pdf.set_font("Times", size=10)
        for term in [t.strip() for t in terms.split(";") if t.strip()]:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 5, f"- {term}")
        pdf.ln(2)

    for line in sanitize_text(text).split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(3)
            continue
        if line.isupper() and len(line) < 100:
            pdf.set_font("Times", "B", 11)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 6, line)
        else:
            pdf.set_font("Times", size=10.5)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 5.2, line)

    return bytes(pdf.output())
